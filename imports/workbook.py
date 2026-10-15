"""
Excel plumbing shared by every sheet: building the downloadable workbook,
reading an uploaded file back into plain row dicts, and marking a file's
problem cells for the admin to fix.

Nothing here knows what a product or an ingredient is — importers.py
describes each sheet with `Column`s and this module turns that description
into a workbook (headers, dropdowns, number checks, a help bubble on every
cell, a "Start here" sheet) and back.
"""

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024
YES_NO = ("Yes", "No")
START_SHEET = "Start here"
LISTS_SHEET = "Lists"

REQUIRED_FILL = PatternFill("solid", fgColor="F4B183")
OPTIONAL_FILL = PatternFill("solid", fgColor="E7E6E6")
ERROR_FILL = PatternFill("solid", fgColor="FF9999")
ESTIMATE_FILL = PatternFill("solid", fgColor="FFF2A8")


@dataclass
class Column:
    key: str
    header: str
    help: str
    required: bool = False
    width: int = 18
    example: str = ""
    # Name of a dropdown list (see build_workbook's `lists`).
    choices: str | None = None
    # strict: Excel refuses anything outside the list. Not strict: the list
    # suggests, and a new value (a new category) asks "keep it?" once.
    strict: bool = False
    # "zero" = number 0 or more, "positive" = number more than 0.
    number: str | None = None

    @property
    def label(self):
        return f"{self.header} *" if self.required else self.header


@dataclass
class SheetSpec:
    title: str
    columns: list[Column]
    rows: list[dict]
    rules: list[str]


@dataclass
class ParsedRow:
    number: int  # the Excel row number the admin sees
    values: dict


@dataclass
class ParsedSheet:
    rows: list[ParsedRow] = field(default_factory=list)
    # column key -> 1-based Excel column, for marking problem cells later.
    positions: dict = field(default_factory=dict)


def _normalise_header(text):
    return " ".join(str(text or "").replace("*", " ").lower().split())


def _range(sheet_title, letter, last_row):
    return f"{quote_sheetname(sheet_title)}!${letter}$2:${letter}${last_row}"


def sheet_column_range(sheet_title, column_index):
    """A whole data column of another sheet, for a dropdown that must offer
    names the admin adds in this same file (recipes using new stock)."""
    return _range(sheet_title, get_column_letter(column_index), MAX_ROWS + 1)


def _validation(column, reference):
    """One validation per column: the dropdown or number rule, plus the help
    bubble Excel shows the moment a cell in the column is selected."""
    if column.choices and reference:
        validation = DataValidation(type="list", formula1=reference, allow_blank=True)
        if column.strict:
            validation.errorStyle = "stop"
            validation.errorTitle = "Not in the list"
            validation.error = f"Pick {column.header} from the dropdown list."
        else:
            validation.errorStyle = "warning"
            validation.errorTitle = "New value"
            validation.error = f"This {column.header.lower()} is not in the list yet. Click Yes to add it as new, or No to fix the spelling."
    elif column.number:
        positive = column.number == "positive"
        validation = DataValidation(type="decimal", operator="greaterThan" if positive else "greaterThanOrEqual", formula1="0", allow_blank=True)
        validation.errorStyle = "stop"
        validation.errorTitle = "Numbers only"
        validation.error = f"{column.header}: write only a number{' more than 0' if positive else ''}, without Rs or units. Example: {column.example or '10'}"
    else:
        validation = DataValidation(allow_blank=True)
    validation.showErrorMessage = bool(column.choices and reference) or bool(column.number)
    validation.showInputMessage = True
    validation.promptTitle = column.label[:32]
    validation.prompt = column.help[:255]
    return validation


def build_workbook(start_lines, sheets, lists):
    """The workbook an admin downloads: a "Start here" sheet, then one sheet
    per SheetSpec pre-filled with current data, then the hidden lists.

    `lists` maps a list name to its dropdown values, or to a ready Excel
    range (str) when the values live on another sheet of the same file."""
    book = Workbook()
    start = book.active
    start.title = START_SHEET

    lists_sheet = book.create_sheet(LISTS_SHEET)
    references = {}
    list_index = 0
    for name, values in lists.items():
        if isinstance(values, str):
            references[name] = values
            continue
        if not values:
            continue
        list_index += 1
        letter = get_column_letter(list_index)
        lists_sheet.cell(row=1, column=list_index, value=name)
        for value_index, value in enumerate(values, start=2):
            lists_sheet.cell(row=value_index, column=list_index, value=value)
        references[name] = _range(LISTS_SHEET, letter, len(values) + 1)
    lists_sheet.sheet_state = "hidden"

    for spec in sheets:
        sheet = book.create_sheet(spec.title, len(book.sheetnames) - 1)
        for index, column in enumerate(spec.columns, start=1):
            letter = get_column_letter(index)
            cell = sheet.cell(row=1, column=index, value=column.label)
            cell.font = Font(bold=True)
            cell.fill = REQUIRED_FILL if column.required else OPTIONAL_FILL
            cell.alignment = Alignment(vertical="center")
            cell.comment = Comment(f"{column.help}\nExample: {column.example}" if column.example else column.help, "Restaurant")
            sheet.column_dimensions[letter].width = column.width
            validation = _validation(column, references.get(column.choices))
            validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
            sheet.add_data_validation(validation)
        sheet.freeze_panes = "A2"
        for row_index, row in enumerate(spec.rows, start=2):
            for col_index, column in enumerate(spec.columns, start=1):
                value = row.get(column.key)
                if value not in (None, ""):
                    sheet.cell(row=row_index, column=col_index, value=value)

    _write_start(start, start_lines, sheets)
    book.active = 0

    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _write_start(sheet, start_lines, sheets):
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 90
    sheet.column_dimensions["D"].width = 22
    row = 1
    for index, line in enumerate(start_lines):
        cell = sheet.cell(row=row, column=1, value=line)
        if index == 0:
            cell.font = Font(bold=True, size=14)
        row += 1
    legend = sheet.cell(row=row + 1, column=1, value="Orange heading = must be filled")
    legend.fill = REQUIRED_FILL
    legend = sheet.cell(row=row + 2, column=1, value="Grey heading = optional")
    legend.fill = OPTIONAL_FILL
    row += 4
    for spec in sheets:
        sheet.cell(row=row, column=1, value=f'Sheet "{spec.title}"').font = Font(bold=True, size=12)
        row += 1
        for rule in spec.rules:
            sheet.cell(row=row, column=1, value=rule)
            row += 1
        for col_index, title in enumerate(("Column", "Must fill?", "What to write", "Example"), start=1):
            sheet.cell(row=row, column=col_index, value=title).font = Font(bold=True)
        row += 1
        for column in spec.columns:
            sheet.cell(row=row, column=1, value=column.header).fill = REQUIRED_FILL if column.required else OPTIONAL_FILL
            sheet.cell(row=row, column=2, value="Yes" if column.required else "No")
            sheet.cell(row=row, column=3, value=column.help).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.cell(row=row, column=4, value=column.example)
            row += 1
        row += 1


def open_upload(upload, *, read_only=True):
    """(workbook, file_errors). Problems that make the whole file unreadable
    come back as file_errors, never raised."""
    if upload is None:
        return None, [{"code": "fileRequired"}]
    if upload.size > MAX_FILE_BYTES:
        return None, [{"code": "fileTooLarge", "params": {"max_mb": MAX_FILE_BYTES // (1024 * 1024)}}]
    try:
        upload.seek(0)
        # data_only: a cell holding a formula is read as the value Excel
        # last calculated, which is what the admin sees on screen.
        return load_workbook(upload, read_only=read_only, data_only=read_only), []
    except Exception:  # openpyxl raises a zoo of types for "not a workbook"
        return None, [{"code": "notExcel"}]


def read_sheet(book, sheet_title, columns):
    """(ParsedSheet, file_errors) for one sheet of an opened workbook."""
    result = ParsedSheet()
    if sheet_title not in book.sheetnames:
        return result, [{"code": "missingSheet", "params": {"sheet": sheet_title}}]
    rows = book[sheet_title].iter_rows(values_only=True)
    header = next(rows, None) or ()
    by_header = {_normalise_header(column.header): column for column in columns}
    for index, text_value in enumerate(header):
        column = by_header.get(_normalise_header(text_value))
        if column is not None and column.key not in result.positions:
            result.positions[column.key] = index + 1
    missing = [column.label for column in columns if column.required and column.key not in result.positions]
    if missing:
        return result, [{"code": "missingColumns", "params": {"sheet": sheet_title, "columns": ", ".join(missing)}}]

    for number, raw in enumerate(rows, start=2):
        values = {key: (raw[index - 1] if index - 1 < len(raw) else None) for key, index in result.positions.items()}
        if all(value is None or str(value).strip() == "" for value in values.values()):
            continue
        if len(result.rows) >= MAX_ROWS:
            return result, [{"code": "tooManyRows", "params": {"sheet": sheet_title, "max": MAX_ROWS}}]
        result.rows.append(ParsedRow(number=number, values=values))
    return result, []


def mark_cells(book, marks, fill=ERROR_FILL):
    """The workbook with every marked cell filled and its message as the
    cell's comment — red problems for the admin to fix where they are, or
    yellow AI estimates for them to check.

    `marks`: [(sheet title, row number, 1-based column, message)]."""
    for sheet_title, row, column, message in marks:
        if sheet_title not in book.sheetnames:
            continue
        cell = book[sheet_title].cell(row=row, column=column)
        cell.fill = fill
        text_value = f"{cell.comment.text}\n{message}" if cell.comment and cell.comment.author == "Check" else message
        cell.comment = Comment(text_value, "Check", width=300, height=110)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


# ── Cell readers. Each returns (value, error) so a row collects every problem
#    at once instead of stopping at the first. ────────────────────────────────

def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # "12" typed into a cell Excel decided was a number
    return " ".join(str(value).split())


def paragraph(value):
    """Free text such as a description: line breaks and spacing are the
    admin's own, so only the edges are trimmed."""
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def decimal(value, column, *, minimum=Decimal("0"), allow_zero=True):
    raw = text(value)
    if raw == "":
        return None, None
    try:
        number = Decimal(raw.replace(",", ""))
    except InvalidOperation:
        return None, {"code": "numberInvalid", "field": column.key, "params": {"column": column.header, "value": raw}}
    if not number.is_finite() or number < minimum or (not allow_zero and number == 0):
        return None, {"code": "numberTooSmall" if allow_zero else "numberPositive", "field": column.key, "params": {"column": column.header, "value": raw}}
    return number, None


def yes_no(value, column):
    raw = text(value).lower()
    if raw == "":
        return None, None
    if raw in ("yes", "y", "true", "1", "نعم"):
        return True, None
    if raw in ("no", "n", "false", "0", "لا"):
        return False, None
    return None, {"code": "yesNoInvalid", "field": column.key, "params": {"column": column.header, "value": text(value)}}


def show(value):
    """How a stored value is echoed back in the preview's old → new."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return YES_NO[0] if value else YES_NO[1]
    if isinstance(value, Decimal):
        return format(value.normalize(), "f")  # 600.00 -> "600", 0.250 -> "0.25"
    return str(value)
