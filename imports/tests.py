import json
from decimal import Decimal
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.test import APIClient

from accounts.models import Staff, StaffRole
from branches.models import Branch
from catalog.models import Category, Product, ProductVariant
from inventory.models import InventoryItem, RecipeItem, StockMovement, StockMovementReason

STOCK = ("Name *", "Unit *", "Current stock", "Par level", "Cost per unit", "Supplier")
MENU = ("Category *", "Name (English)", "Name (Arabic)", "Description (English)", "Description (Arabic)",
        "Price", "Available", "Size (English)", "Size (Arabic)", "Size price", "Recipe multiplier")
RECIPES = ("Dish *", "Ingredient *", "Quantity *", "Unit *")


def make_staff(branch, name, role):
    staff = Staff(branch=branch, name=name, role=role, is_active=True)
    staff.set_pin("1234")
    staff.save()
    return staff


def xlsx(stock=(), menu=(), recipes=(), *, sheets=None):
    """The restaurant workbook. Each sheet gets its header row even when it
    has no rows, like a template the admin only changed one sheet of."""
    book = Workbook()
    book.remove(book.active)
    for title, headers, rows in sheets or (("1 Stock", STOCK, stock), ("2 Menu", MENU, menu), ("3 Recipes", RECIPES, recipes)):
        sheet = book.create_sheet(title)
        sheet.append(list(headers))
        for row in rows:
            sheet.append(list(row))
    output = BytesIO()
    book.save(output)
    return SimpleUploadedFile("import.xlsx", output.getvalue())


class ImportTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.branch = Branch.objects.create(name_en="Farzan", name_ar="فرزان", timezone="Asia/Karachi", currency="PKR")
        cls.owner = make_staff(cls.branch, "Owner", StaffRole.OWNER)
        cls.cashier = make_staff(cls.branch, "Cashier", StaffRole.CASHIER)
        cls.chicken = InventoryItem.objects.create(branch=cls.branch, name="Chicken", unit="kg", current_stock=10, par_level=2, cost_per_unit=600)
        cls.category = Category.objects.create(branch=cls.branch, name_en="Karahi", name_ar="")
        cls.karahi = Product.objects.create(branch=cls.branch, category=cls.category, name_en="Chicken Karahi", name_ar="", price=1400, cost_price=0)
        cls.half = ProductVariant.objects.create(product=cls.karahi, name_en="Half KG", name_ar="", price=1400, recipe_multiplier=1)

    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(user=self.owner)

    def upload(self, file, commit=False):
        data = {"file": file}
        if commit:
            data["commit"] = "1"
        return self.client.post("/api/imports/", data, format="multipart")

    def rows(self, body, key):
        return next(sheet for sheet in body["sheets"] if sheet["key"] == key)["rows"]


class TemplateTests(ImportTestCase):
    def test_template_has_start_sheet_then_prefilled_sheets_in_order(self):
        response = self.client.get("/api/imports/template/")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames[:4], ["Start here", "1 Stock", "2 Menu", "3 Recipes"])
        stock = book["1 Stock"]
        self.assertEqual((stock["A1"].value, stock["A2"].value, stock["B2"].value), ("Name *", "Chicken", "kg"))
        self.assertEqual(book["2 Menu"]["B2"].value, "Chicken Karahi")

    def test_cells_explain_themselves_and_refuse_wrong_values(self):
        book = load_workbook(BytesIO(self.client.get("/api/imports/template/").content))
        stock = book["1 Stock"]
        by_column = {str(v.sqref).split(":")[0][0]: v for v in stock.data_validations.dataValidation}
        unit, cost = by_column["B"], by_column["E"]
        self.assertEqual((unit.type, unit.errorStyle), ("list", "stop"))
        self.assertEqual((cost.type, cost.operator, cost.errorStyle), ("decimal", "greaterThanOrEqual", "stop"))
        self.assertTrue(cost.showInputMessage)
        self.assertIn("ONE unit", cost.prompt)
        recipes = {str(v.sqref).split(":")[0][0]: v for v in book["3 Recipes"].data_validations.dataValidation}
        # Ingredients come straight from the Stock sheet of the same file.
        self.assertEqual(recipes["B"].formula1, "'1 Stock'!$A$2:$A$5001")
        self.assertEqual(recipes["B"].errorStyle, "stop")

    def test_only_owner_or_manager_can_import(self):
        self.client.force_authenticate(user=self.cashier)
        self.assertEqual(self.client.get("/api/imports/template/").status_code, status.HTTP_403_FORBIDDEN)


class FileTests(ImportTestCase):
    def test_not_an_excel_file(self):
        body = self.upload(SimpleUploadedFile("x.xlsx", b"hello")).json()
        self.assertEqual(body["file_errors"], [{"code": "notExcel"}])

    def test_missing_sheet_and_missing_columns(self):
        body = self.upload(xlsx(sheets=(("1 Stock", ("Name",), [("Oil",)]), ("2 Menu", MENU, []),))).json()
        codes = [error["code"] for error in body["file_errors"]]
        self.assertEqual(codes, ["missingColumns", "missingSheet"])
        self.assertEqual(body["file_errors"][1]["params"], {"sheet": "3 Recipes"})

    def test_empty_file(self):
        self.assertEqual(self.upload(xlsx()).json()["file_errors"], [{"code": "noRows"}])


class StockImportTests(ImportTestCase):
    def test_preview_shows_changes_without_writing(self):
        body = self.upload(xlsx(stock=[("chicken", "kg", 12, None, 650, None), ("Oil", "l", 5, 1, 400, "Al Noor")])).json()
        self.assertEqual((body["summary"]["update"], body["summary"]["create"]), (1, 1))
        update, create = self.rows(body, "stock")
        self.assertIn({"field": "cost_per_unit", "old": "600", "new": "650"}, update["changes"])
        self.assertIn({"field": "current_stock", "old": "10", "new": "12"}, update["changes"])
        self.assertEqual(create["notes"], [{"code": "newSupplier", "params": {"name": "Al Noor"}}])
        self.chicken.refresh_from_db()
        self.assertEqual(self.chicken.cost_per_unit, Decimal("600.00"))
        self.assertFalse(InventoryItem.objects.filter(name="Oil").exists())

    def test_commit_applies_and_records_stock_adjustments(self):
        response = self.upload(xlsx(stock=[("Chicken", "kg", 12, None, 650, None), ("Oil", "l", 5, 1, 400, "Al Noor")]), commit=True)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertTrue(response.json()["committed"])
        self.chicken.refresh_from_db()
        self.assertEqual((self.chicken.current_stock, self.chicken.cost_per_unit), (Decimal("12"), Decimal("650")))
        movement = StockMovement.objects.get(inventory_item=self.chicken)
        self.assertEqual((movement.quantity_delta, movement.reason), (Decimal("2"), StockMovementReason.ADJUSTMENT))
        oil = InventoryItem.objects.get(branch=self.branch, name="Oil")
        self.assertEqual((oil.current_stock, oil.supplier.name_en), (Decimal("5"), "Al Noor"))

    def test_any_error_on_any_sheet_blocks_the_whole_commit(self):
        file = xlsx(stock=[("Oil", "l", 5)], menu=[("Karahi", "Lassi", None, None, None, "three hundred")])
        response = self.upload(file, commit=True)
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        error = self.rows(response.json(), "products")[0]["errors"][0]
        self.assertEqual((error["code"], error["field"]), ("numberInvalid", "price"))
        self.assertFalse(InventoryItem.objects.filter(name="Oil").exists())

    def test_duplicate_names_are_reported(self):
        rows = self.rows(self.upload(xlsx(stock=[("Oil", "l"), ("oil", "l")])).json(), "stock")
        self.assertEqual(rows[1]["errors"][0], {"code": "duplicateInFile", "field": "name", "params": {"firstRow": 2}})

    def test_new_item_close_to_an_existing_one_is_flagged(self):
        row = self.rows(self.upload(xlsx(stock=[("Chiken", "kg")])).json(), "stock")[0]
        self.assertEqual(row["action"], "create")
        self.assertIn({"code": "similarExists", "params": {"name": "Chiken", "existing": "Chicken"}}, row["notes"])

    def test_unit_change_that_breaks_a_recipe_is_rejected(self):
        RecipeItem.objects.create(product=self.karahi, inventory_item=self.chicken, quantity=500, unit="g")
        row = self.rows(self.upload(xlsx(stock=[("Chicken", "pcs")])).json(), "stock")[0]
        self.assertEqual(row["errors"][0]["code"], "unitBreaksRecipes")

    def test_unit_change_is_fine_when_the_file_also_fixes_the_recipe(self):
        RecipeItem.objects.create(product=self.karahi, inventory_item=self.chicken, quantity=500, unit="g")
        body = self.upload(xlsx(stock=[("Chicken", "pcs")], recipes=[("Chicken Karahi", "Chicken", 2, "pcs")])).json()
        self.assertEqual(body["summary"]["error"], 0)


class ProductImportTests(ImportTestCase):
    def test_new_size_on_existing_dish_and_new_dish_in_new_category(self):
        rows = [
            ("Karahi", "Chicken Karahi", None, None, None, 1400, "Yes", "Half KG", None, 1400, 1),
            ("Karahi", "Chicken Karahi", None, None, None, 1400, "Yes", "1 KG", None, 2600, 2),
            ("Drinks", "Lassi", None, None, None, 300, "No", None, None, None, None),
        ]
        body = self.upload(xlsx(menu=rows)).json()
        results = self.rows(body, "products")
        self.assertEqual([row["action"] for row in results], ["unchanged", "create", "create"])
        self.assertEqual(results[2]["notes"], [{"code": "newCategory", "params": {"name": "Drinks"}}])

        self.assertEqual(self.upload(xlsx(menu=rows), commit=True).status_code, status.HTTP_200_OK)
        full = ProductVariant.objects.get(product=self.karahi, name_en="1 KG")
        self.assertEqual((full.price, full.recipe_multiplier), (Decimal("2600"), Decimal("2")))
        lassi = Product.objects.get(branch=self.branch, name_en="Lassi")
        self.assertEqual((lassi.category.name_en, lassi.is_available), ("Drinks", False))

    def test_price_change_is_previewed_old_to_new(self):
        row = self.rows(self.upload(xlsx(menu=[("Karahi", "Chicken Karahi", None, None, None, 1500)])).json(), "products")[0]
        self.assertIn({"field": "price", "old": "1400", "new": "1500"}, row["changes"])

    def test_conflicting_dish_values_across_size_rows(self):
        menu = [
            ("Karahi", "Chicken Karahi", None, None, None, 1400, None, "Half KG", None, 1400, 1),
            ("Karahi", "Chicken Karahi", None, None, None, 1500, None, "1 KG", None, 2600, 2),
        ]
        row = self.rows(self.upload(xlsx(menu=menu)).json(), "products")[1]
        self.assertEqual(row["errors"][0], {"code": "conflictingValue", "field": "price", "params": {"column": "Price", "firstRow": 2}})

    def test_new_dish_needs_a_price(self):
        row = self.rows(self.upload(xlsx(menu=[("Karahi", "Mutton Karahi")])).json(), "products")[0]
        self.assertEqual(row["errors"][0]["code"], "priceRequired")


class RecipeImportTests(ImportTestCase):
    def setUp(self):
        super().setUp()
        self.oil = InventoryItem.objects.create(branch=self.branch, name="Oil", unit="l", cost_per_unit=400)
        self.tomato = InventoryItem.objects.create(branch=self.branch, name="Tomato", unit="kg", cost_per_unit=200)
        RecipeItem.objects.create(product=self.karahi, inventory_item=self.chicken, quantity=500, unit="g")
        RecipeItem.objects.create(product=self.karahi, inventory_item=self.tomato, quantity=200, unit="g")

    def test_recipe_is_replaced_and_removed_lines_are_previewed(self):
        recipes = [("Chicken Karahi", "Chicken", 600, "g"), ("Chicken Karahi", "Oil", 50, "ml")]
        rows = self.rows(self.upload(xlsx(recipes=recipes)).json(), "recipes")
        self.assertEqual([row["action"] for row in rows], ["update", "create", "delete"])
        self.assertEqual(rows[2]["name"], "Chicken Karahi · Tomato")

        self.assertEqual(self.upload(xlsx(recipes=recipes), commit=True).status_code, status.HTTP_200_OK)
        lines = {line.inventory_item.name: (line.quantity, line.unit) for line in RecipeItem.objects.filter(product=self.karahi)}
        self.assertEqual(lines, {"Chicken": (Decimal("600"), "g"), "Oil": (Decimal("50"), "ml")})
        self.karahi.refresh_from_db()
        # 0.6 kg x 600 + 0.05 l x 400 = 360 + 20
        self.assertEqual(self.karahi.cost_price, Decimal("380.00"))

    def test_typos_get_a_suggestion_instead_of_a_new_item(self):
        recipes = [("Biryani", "Chicken", 1, "kg"), ("Chiken Karahi", "Chicken", 1, "kg"), ("Chicken Karahi", "Tomatto", 1, "g"), ("Chicken Karahi", "Chicken", 2, "pcs")]
        rows = self.rows(self.upload(xlsx(recipes=recipes)).json(), "recipes")
        self.assertEqual([row["errors"][0]["code"] for row in rows[:4]], ["dishNotFound", "dishSimilar", "ingredientSimilar", "unitIncompatible"])
        self.assertEqual(rows[2]["errors"][0]["params"], {"name": "Tomatto", "existing": "Tomato"})
        self.assertFalse(InventoryItem.objects.filter(name="Tomatto").exists())

    def test_new_dish_and_new_stock_from_the_same_file(self):
        def file():
            return xlsx(
                stock=[("Rice", "kg", 20, None, 300)],
                menu=[("Rice", "Chicken Pulao", None, None, None, 900)],
                recipes=[("Chicken Pulao", "Rice", 250, "g"), ("Chicken Pulao", "Chicken", 200, "g")],
            )
        body = self.upload(file()).json()
        self.assertEqual(body["summary"]["error"], 0)
        self.assertEqual(self.upload(file(), commit=True).status_code, status.HTTP_200_OK)
        pulao = Product.objects.get(branch=self.branch, name_en="Chicken Pulao")
        self.assertEqual({line.inventory_item.name for line in pulao.recipe_items.all()}, {"Rice", "Chicken"})
        # 0.25 kg x 300 + 0.2 kg x 600
        self.assertEqual(pulao.cost_price, Decimal("195.00"))

    def test_ingredient_only_in_a_recipe_is_added_to_stock_visibly(self):
        recipes = [("Chicken Karahi", "Chicken", 500, "g"), ("Chicken Karahi", "Tomato", 200, "g"), ("Chicken Karahi", "Zeera", 5, "g")]
        row = self.rows(self.upload(xlsx(recipes=recipes)).json(), "recipes")[2]
        self.assertEqual((row["action"], row["notes"]), ("create", [{"code": "newIngredient", "params": {"name": "Zeera", "unit": "g"}}]))
        self.assertEqual(self.upload(xlsx(recipes=recipes), commit=True).status_code, status.HTTP_200_OK)
        zeera = InventoryItem.objects.get(branch=self.branch, name="Zeera")
        self.assertEqual((zeera.unit, zeera.current_stock, zeera.cost_per_unit), ("g", Decimal("0"), Decimal("0")))


class GuideTests(ImportTestCase):
    def test_guide_explains_columns_and_counts_setup(self):
        body = self.client.get("/api/imports/guide/").json()
        self.assertEqual([sheet["title"] for sheet in body["sheets"]], ["1 Stock", "2 Menu", "3 Recipes"])
        unit = next(column for column in body["sheets"][0]["columns"] if column["key"] == "unit")
        self.assertTrue(unit["required"])
        self.assertEqual(body["checklist"], {"stock_items": 1, "dishes": 1, "dishes_without_recipe": 1})
        self.assertIn("Dish * | Ingredient * | Quantity * | Unit *", body["prompt"])
        self.assertIn("Chicken Karahi", body["prompt"])


class MarkedFileTests(ImportTestCase):
    def test_problem_cells_come_back_red_with_the_message(self):
        file = xlsx(stock=[("Oil", "l", "lots")])
        marks = [
            {"sheet": "stock", "row": 2, "field": "current_stock", "text": "Current stock: write a number"},
            {"sheet": "stock", "row": 2, "field": None, "text": "Whole row problem"},
            {"sheet": "nope", "row": 2, "field": None, "text": "ignored"},
        ]
        response = self.client.post("/api/imports/marked/", {"file": file, "marks": json.dumps(marks)}, format="multipart")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        sheet = load_workbook(BytesIO(response.content))["1 Stock"]
        self.assertEqual(sheet["C2"].comment.text, "Current stock: write a number")
        self.assertEqual(sheet["C2"].fill.fgColor.rgb[-6:], "FF9999")
        self.assertEqual(sheet["A2"].comment.text, "Whole row problem")
